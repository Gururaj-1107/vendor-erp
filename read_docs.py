
import sys
import os
import codecs

sys.stdout = codecs.getwriter('utf-8')(sys.stdout.detach())

def read_docx(path):
    try:
        from docx import Document
        doc = Document(path)
        print(f"\n\n{'='*80}")
        print(f"FILE: {path}")
        print('='*80)
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = '\t'.join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    print(row_text)
    except Exception as e:
        print(f"Error reading {path}: {e}")

def read_pdf(path):
    try:
        import pdfplumber
        print(f"\n\n{'='*80}")
        print(f"FILE: {path}")
        print('='*80)
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    print(text)
    except ImportError:
        try:
            import fitz  # PyMuPDF
            print(f"\n\n{'='*80}")
            print(f"FILE: {path}")
            print('='*80)
            doc = fitz.open(path)
            for page in doc:
                print(page.get_text())
        except Exception as e:
            print(f"Error reading PDF {path}: {e}")

os.system("pip install python-docx pdfplumber 2>nul")

read_docx(r"d:\Vendor\VendorBridge Procurement.docx")
read_docx(r"d:\Vendor\webpages_details.docx")
read_pdf(r"d:\Vendor\Vendorbridge Hackathon Problem Statement.pdf")
read_docx(r"d:\Vendor\frontend_codes.docx")
